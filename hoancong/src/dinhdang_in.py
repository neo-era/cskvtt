# -*- coding: utf-8 -*-
"""Chinh dinh dang in toan bo file ho so: A4, vua 1 trang ngang (Bia 1x1), vung in den het chu ky,
   lap tieu de bang, noi chieu cao dong chu dai (chi tang, bo dong dem <10pt, bo BIA).  python3 dinhdang_in.py IN.xlsx OUT.xlsx"""
import openpyxl, warnings, sys, re, math; warnings.filterwarnings('ignore')
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.cell.cell import MergedCell
def wrapn(p,cpl):            # dem so dong theo ngat TU (khong cat giua tu)
    n,cur=1,0
    for w in p.split(' '):
        L=len(w)
        if cur==0: cur=L
        elif cur+1+L<=cpl: cur+=1+L
        else: n+=1; cur=L
        while cur>cpl: n+=1; cur-=cpl
    return n
K=0.42          # he so do rong chu Times (hieu chinh tu cac dong mau da duyet)
TITLES={'4_G201_TTH_PTH_TMY_THU':'$4:$5'}
def fmt(src,out):
    print('@@Định dạng in')
    wb=openpyxl.load_workbook(src); log=[]
    for s in wb.worksheets:
        if s.sheet_state!='visible': continue
        pa=s.print_area; c1,c2=1,s.max_column
        if pa:
            m=re.search(r"\$?([A-Z]+)\$?\d*:\$?([A-Z]+)\$?\d*",(pa if isinstance(pa,str) else pa[0]).split(',')[0].split('!')[-1])
            if m: c1,c2=CI(m.group(1)),CI(m.group(2))
        last=max([r for r in range(1,s.max_row+1) if any(s.cell(r,c).value not in (None,'') for c in range(c1,c2+1))]+[1])
        s.print_area='%s1:%s%d'%(L(c1),L(c2),last+1)
        ps=s.page_setup; ps.paperSize=9; ps.scale=None
        s.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
        ps.fitToWidth=1; ps.fitToHeight=1 if s.title.startswith('BIA') else 0
        if s.title in TITLES and not s.print_title_rows: s.print_title_rows=TITLES[s.title]
        if s.title.startswith('BIA'): log.append((s.title,s.print_area,0)); continue
        span={(m.min_row,m.min_col):(m.max_row-m.min_row+1,m.min_col,m.max_col) for m in s.merged_cells.ranges}
        WD={}                                   # do rong tung cot (openpyxl gop cot lien nhau cung do rong vao 1 muc min..max)
        for d in s.column_dimensions.values():
            for c in range(d.min or 1,(d.max or d.min or 1)+1): WD[c]=(0 if d.hidden else (d.width or 9.14))
        def cw(c): return WD.get(c,9.14)
        ch=0
        for row in s.iter_rows(max_row=last):
            r=row[0].row; cur=s.row_dimensions[r].height
            if cur is not None and cur<10: continue                      # dong dem cua mau
            need=0
            for x in row:
                if isinstance(x,MergedCell) or x.value is None or isinstance(x.value,(int,float)): continue
                v=str(x.value)
                if v.startswith('='): continue
                nr,a,b=span.get((r,x.column),(1,x.column,x.column))
                if nr>1: continue
                bd=x.border
                if not (bd.left.style or bd.right.style or bd.top.style or bd.bottom.style) and b-a<2: continue   # chi o bang hoac khoi doan van
                wpx=sum(cw(c)*7+5 for c in range(a,b+1))-6; sz=x.font.sz or 11
                cpl=max(1,int(wpx/(K*sz*1.333*(1.08 if x.font.b else 1))*0.97))   # chu dam rong hon ~8%
                paras=v.split('\n')
                lines=sum(wrapn(p,cpl) for p in paras) if x.alignment.wrap_text else len(paras)
                need=max(need,lines*sz*1.27+2)
            if need>(cur or 15.75)+1: s.row_dimensions[r].height=round(need*4)/4; ch+=1
        log.append((s.title,s.print_area,ch))
    wb.save(out)
    return log
if __name__=='__main__':
    for l in fmt(sys.argv[1],sys.argv[2]): print('%-28s vùng in %-40s nới %d dòng'%l)
