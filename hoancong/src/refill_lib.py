# -*- coding: utf-8 -*-
import re, unicodedata, glob, os, copy
import pdfplumber, openpyxl
D=os.environ.get('SRC_DIR','./')            # thu muc PDF nguon (co / cuoi)
MASTER=os.environ['MASTER_XLSX']           # file Control_Cabinet
QUAN=os.environ.get('QUAN','8')            # ma quan trong bien ban chi tiet
def sa(s):
    s=unicodedata.normalize('NFD',str(s or '')); s=''.join(c for c in s if unicodedata.category(c)!='Mn')
    return s.replace('đ','d').replace('Đ','D')
def norm(s): return re.sub(r'\s+',' ',re.sub(r'[^a-z0-9]+',' ',sa(s).lower())).strip()
def norm_ns(s): return re.sub(r'[^a-z0-9]','',sa(s).lower())
def num(s):
    if s is None: return None
    s=str(s).strip().replace('.','').replace(',','.')
    try: return float(s)
    except: return None
from difflib import get_close_matches
GOI=None            # loc theo goi thau (vd "G201NA"); None = khong loc
PREFER=set()         # uu tien phuong trong tap nay khi con nhieu ung vien
def master_rows():
    """(ten_tu, phuong, duong, diachi, goi) tu master .json (ban gon) hoac .xlsx Control_Cabinet."""
    if MASTER.lower().endswith('.json'):
        import json
        return [tuple(r) for r in json.load(open(MASTER,encoding='utf-8'))]
    wb=openpyxl.load_workbook(MASTER,data_only=True,read_only=True)
    if 'BANG_TRA_TU' in wb.sheetnames:   # file chuan: B ten tu, C phuong, E duong, F dia chi, G goi
        return [(r[1],r[2],r[4],r[5],r[6]) for i,r in enumerate(wb['BANG_TRA_TU'].iter_rows(values_only=True)) if i and r[1]]
    ws=wb['Control_Cabinet_cs(tudieukhien)']
    out=[]
    for i,r in enumerate(ws.iter_rows(values_only=True)):
        if i==0 or not r[7]: continue
        out.append((r[7],r[2],r[5],r[8],r[11]))
    return out
def load_master():
    by,bn={},{}
    for ten,ph,du,dc,goi in master_rows():
        rec=(ph,du,dc,goi,ten)  # phuong,duong,diachi,goi,ten tu chuan
        by.setdefault(norm(ten),[]).append(rec)
        bn.setdefault(norm_ns(ten),[]).append(rec)
    return by,bn
def _pick(cands):
    if not cands: return None
    c=cands
    if GOI: 
        g=[x for x in c if x[3]==GOI]
        if g: c=g
    if len(c)>1 and PREFER:
        p=[x for x in c if x[0] in PREFER]
        if p: c=p
    return (c[0][0],c[0][1],c[0][2],c[0][4])
def look(t,by,bn):
    k=norm(t)
    if k in by: return _pick(by[k])
    kn=norm_ns(t)
    if kn in bn: return _pick(bn[kn])
    m=get_close_matches(k,list(by),n=1,cutoff=0.9); return _pick(by[m[0]]) if m else None
_FILES=None
def detect_files():
    """Nhan dien PDF theo NOI DUNG (ten file co the dat lech), dua vao do rong bang o 2 trang dau:
       bang 11 cot = BD chi tiet; 9 cot = QLVH chi tiet; 7 cot >=25 dong = BoQ 499."""
    global _FILES
    if _FILES: return _FILES
    from collections import Counter
    res={}
    for p in sorted(glob.glob(D+'*.pdf')+glob.glob(D+'*.PDF')):
        c=Counter()
        with pdfplumber.open(p) as pdf:
            for pg in pdf.pages[:2]:
                for tb in pg.extract_tables():
                    for r in tb: c[len(r)]+=1
        if c[11]>=5: res.setdefault('BD',p)
        elif c[9]>=5: res.setdefault('QLVH',p)
        elif c[7]>=25: res.setdefault('BOQ',p)
    miss=[k for k in ('QLVH','BOQ','BD') if k not in res]
    if miss: raise RuntimeError('Không tìm thấy biên bản: '+', '.join({'QLVH':'Xác nhận KL QLVH chi tiết (bảng 9 cột)','BOQ':'Nghiệm thu KL bảo dưỡng 499 hạng mục','BD':'Xác nhận KL bảo dưỡng chi tiết (bảng 11 cột)'}[k] for k in miss))
    _FILES=res; return res
_KEY={'4._BBXNKLHT_QLVH':'QLVH','2._BBNT_QLVH':'BOQ','1._BBNTHT_QLVH':'BD'}
def g(k): return detect_files()[_KEY[k]]
_RNG=[('1500m3000m','1500m - 3000m'),('1000m1500m','1000m - 1500m'),('500m1000m','500m - 1000m')]
def canon_hm(a,b=''):
    """Ten hang muc QLVH trong PDF bi cat sang cot ben ('... (tron' + 'g hem) - ket noi...').
    Ghep lai va dung ten chuan: 'Duy tri tram den <khoang>[ (trong hem)][ (khu dan cu)][ - ket noi ve trung tam dieu khien]'."""
    raw=(a or '')+(b or ''); k=norm_ns(raw); pre='duytritramden'
    if not k.startswith(pre): return ((a or '')+' '+(b or '')).strip()
    rest=k[len(pre):]; lab=None
    for key,l in _RNG:
        if rest.startswith(key): lab=l; rest=rest[len(key):]; break
    if lab is None:
        for key,sign in (('3000m','>'),('500m','<')):
            if rest.startswith(key) and sign in raw: lab=sign+' '+key; rest=rest[len(key):]; break
    if lab is None: return ((a or '')+' '+(b or '')).strip()
    txt='Duy trì trạm đèn '+lab
    if rest.startswith('tronghem'): txt+=' (trong hẻm)'; rest=rest[8:]
    if rest.startswith('khudancu'): txt+=' (khu dân cư)'; rest=rest[8:]
    if rest.startswith('ketnoivetrungtamdieukhien'): txt+=' - kết nối về trung tâm điều khiển'; rest=rest[25:]
    return txt if rest=='' else ((a or '')+' '+(b or '')).strip()
def parse_qlvh(by,bn):
    out=[]; hm=None; grp=None
    with pdfplumber.open(g('4._BBXNKLHT_QLVH')) as pdf:
        for pg in pdf.pages:
            for tb in pg.extract_tables():
                for r in tb:
                    r=[(c or '').strip().replace('\n',' ') for c in r]
                    if len(r)<9: continue
                    c0,c1,c2,c3,c4,c5,c6,c7,c8=r[:9]
                    if c1 in ('Tủ Thường','Tủ kết nối về trung tâm') and not c4: grp=c1; continue
                    if c1.startswith('Tổng cộng'): continue
                    if c0 and c1 and not c3 and not c4: hm=canon_hm(c1,c2); continue
                    if not c0 and c1 and not c3 and not c4: continue
                    if c3==QUAN:
                        m=look(c1,by,bn); ph=m[0] if m else '(không rõ)'; du=m[1] if m else ''
                        out.append({'tu':(m[3] if m else c1),'diachi':c2,'kl':num(c6),'hm':hm,'grp':grp,'phuong':ph,'duong':du})
    return out
def parse_boq():
    it={}
    with pdfplumber.open(g('2._BBNT_QLVH')) as pdf:
        for pg in pdf.pages:
            for tb in pg.extract_tables():
                for r in tb:
                    r=[(c or '').strip().replace('\n',' ') for c in r]
                    if len(r)<6: continue
                    st=r[0].strip()
                    if re.fullmatch(r'\d+',st): it[int(st)]={'ten':r[1],'dvt':r[2],'kh':num(r[3]),'th':num(r[4]),'nt':num(r[5])}
    return it
def parse_bd(by,bn):
    out=[]; hm=None; dvt=None
    with pdfplumber.open(g('1._BBNTHT_QLVH')) as pdf:
        for pg in pdf.pages:
            for tb in pg.extract_tables():
                for r in tb:
                    r=[(c or '').strip().replace('\n',' ') for c in r]
                    if len(r)<11: continue
                    st,ten,u,quan,duong,tentu,mota,kh,th,nt,gc=r[:11]
                    if re.fullmatch(r'\d+',st.strip()) and ten and not tentu: hm=ten; dvt=u; continue
                    if quan==QUAN and tentu:
                        m=look(tentu,by,bn); ph=m[0] if m else '(không rõ)'
                        out.append({'hm':hm,'dvt':dvt,'duong':duong,'tu':(m[3] if m else tentu),'mota':mota,'kl':num(nt),'phuong':ph})
    return out

def rewrite(ws, first_data, tail_start, ncols, nnew, writer, probe_row):
    """Capture tail (tail_start..max_row), clear data+tail, write nnew data rows via writer, re-place tail."""
    maxr=ws.max_row
    # capture tail cells
    tail=[]
    for r in range(tail_start, maxr+1):
        for c in range(1, ncols+1):
            cell=ws.cell(r,c)
            tail.append((r-tail_start, c, cell.value, copy.copy(cell._style) if cell.has_style else None))
    # capture tail merges + heights
    tail_merges=[]
    for m in list(ws.merged_cells.ranges):
        if m.min_row>=tail_start:
            tail_merges.append((m.min_row-tail_start,m.max_row-tail_start,m.min_col,m.max_col))
    heights={ r-tail_start: ws.row_dimensions[r].height for r in range(tail_start,maxr+1) if ws.row_dimensions[r].height }
    # remove all merges from first_data down (proper unmerge)
    for m in list(ws.merged_cells.ranges):
        if m.max_row>=first_data: ws.unmerge_cells(str(m))
    # clear cells first_data..maxr
    from openpyxl.styles.cell_style import StyleArray
    for r in range(first_data, maxr+1):
        ws.row_dimensions[r].height=None
        for c in range(1, ncols+1):
            cell=ws.cell(r,c)
            if not isinstance(cell, openpyxl.cell.cell.MergedCell):
                cell.value=None; cell._style=StyleArray()   # xoa sach style cu -> khong con duong ke thua
    # write data
    data_merges=writer(ws, first_data, probe_row)
    new_tail=first_data+nnew
    # place tail
    for roff,c,val,st in tail:
        cell=ws.cell(new_tail+roff, c); cell.value=val
        if st is not None: cell._style=st
    for a,b,c1,c2 in tail_merges:
        ws.merge_cells(start_row=new_tail+a,end_row=new_tail+b,start_column=c1,end_column=c2)
    for roff,h in heights.items(): ws.row_dimensions[new_tail+roff].height=h
    for (r1,r2,c1,c2) in data_merges:
        ws.merge_cells(start_row=r1,end_row=r2,start_column=c1,end_column=c2)
    return new_tail