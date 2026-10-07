# -*- coding: utf-8 -*-
import warnings, copy, re; warnings.filterwarnings('ignore')
import openpyxl
from refill_lib import *
from collections import OrderedDict
import os
SRC=os.environ['SRC_XLSM']; OUT=os.environ.get('OUT_XLSX','HO_SO_HOAN_CONG.xlsx')
import refill_lib as _RL
_RL.GOI=os.environ.get('GOI','G201NA')
_RL.PREFER=set(p.strip() for p in os.environ.get('PREFER','').split(';') if p.strip())
print('@@Đọc bảng tra tủ → phường')
by,bn=load_master()
print('@@Đọc biên bản PDF (bước lâu nhất)')
qlvh=parse_qlvh(by,bn); boq=parse_boq(); bd=parse_bd(by,bn)
# PDF mat dau 'ó' (bong->bóng): lay ten hang muc + DVT chuan tu chinh file mau (6.HTCT_BD) theo ten khong dau
print('@@Mở file mẫu')
wb=openpyxl.load_workbook(SRC, data_only=False, keep_vba=False)   # bo macro (load 1 lan)
# Tu nhan dien ky/quan cua FILE MAU neu khong truyen (BIA 'THUC HIEN THANG x', so hieu '..-x-Qy')
if not os.environ.get('THANG_MAU'):
    for _row in wb['BIA_G201'].iter_rows():
        for _c in _row:
            _mm=re.search(r'TH[ÁA]NG\s+0?(\d{1,2})\s+N[ĂA]M',str(_c.value or '')) if isinstance(_c.value,str) else None
            if _mm and not os.environ.get('THANG_MAU'): os.environ['THANG_MAU']=_mm.group(1)
if not os.environ.get('QUAN_MAU'):
    for _s in wb.worksheets:
        for _row in _s.iter_rows(max_row=12):
            for _c in _row:
                _mm=re.search(r'(?:YCNT|NTCV)-\d{1,2}-Q(\d+)',_c.value) if isinstance(_c.value,str) else None
                if _mm and not os.environ.get('QUAN_MAU'): os.environ['QUAN_MAU']=_mm.group(1)
print('@@Mẫu: kỳ tháng %s, quận %s'%(os.environ.get('THANG_MAU'),os.environ.get('QUAN_MAU')))
DB_MAU_TXT=next((c.value for row in wb['BIA_G201'].iter_rows() for c in row if isinstance(c.value,str) and c.value.strip().startswith('Địa bàn')),'')
_t=wb['6.HTCT_BD_G201']; TN={}
for _r in range(33,_t.max_row+1):
    _n=_t.cell(_r,2).value
    if isinstance(_n,str) and _t.cell(_r,4).value: TN[norm(_n)]=(_n,_t.cell(_r,4).value)
for _it in boq.values():
    if norm(_it['ten']) in TN: _it['ten'],_it['dvt']=TN[norm(_it['ten'])]
for _d in bd:
    if norm(_d['hm'] or '') in TN: _d['hm'],_d['dvt']=TN[norm(_d['hm'])]
# ===== VUNG BAO VE: can cu phap ly (lay tu MAU, so theo gia tri) -> khong doi bat ky chu nao =====
_H=re.compile(r'^\s*(\d+|[IVX]+)\s*[\)\.]\s*\S')
PROTECT=set()
for _s in wb.worksheets:
    st=None
    for r in range(1,min(_s.max_row,90)+1):
        vals=[_s.cell(r,c).value for c in range(1,4) if isinstance(_s.cell(r,c).value,str)]
        t=' '.join(vals); tl=t.lower().strip()
        if st is None:
            if 'căn cứ' in tl or tl.startswith('- nghị định'): st=r
            else: continue
        elif _H.match(t) or tl.startswith('- nội dung') or tl.startswith('stt'): break
        PROTECT.update(vals)
    for row in _s.iter_rows():
        for c in row:
            if isinstance(c.value,str) and 'Căn cứ' in c.value: PROTECT.add(c.value)
from collections import Counter as _C
_c=_C()
for _d in qlvh+bd: _c[_d['phuong']]+=1
DB=[p for p,n in _c.most_common() if n>=5 and p!='(không rõ)']
DBSTR=', '.join(DB)

def markers(ws):
    tong=sig=None
    for r in range(1,ws.max_row+1):
        for c in range(1,14):
            v=ws.cell(r,c).value
            if isinstance(v,str):
                if tong is None and v.strip().startswith('Tổng cộng'): tong=r
                if sig is None and 'ĐẠI DIỆN' in v: sig=r
    return tong,sig
def find_tail(ws, after):
    keys=('Thời gian','Chỉ huy','ĐẠI DIỆN','Đại diện','Kính gửi','KẾT LUẬN','Kết luận')
    for r in range(after+1, ws.max_row+1):
        for c in range(1,6):
            v=ws.cell(r,c).value
            if isinstance(v,str) and any(k in v for k in keys): return r
    return ws.max_row+1
def capstyle(ws,row,ncol):
    return {c:(copy.copy(ws.cell(row,c)._style) if ws.cell(row,c).has_style else None) for c in range(1,ncol+1)}
def caph(ws,r): return ws.row_dimensions[r].height
def seth(ws,r,h): ws.row_dimensions[r].height=h
WROTE={}  # sheet -> (r1,r2) vung du lieu da ghi (de autofit chi vung nay)
def sphuong(p): return re.sub(r'^(Phường|Xã)\s+','',p or '')
def putstyle(ws,row,st,ncol):
    for c in range(1,ncol+1):
        if st.get(c) is not None: ws.cell(row,c)._style=st[c]

# HM order QLVH
# Nhan nhom theo DUNG CHU CUA MAU (3.HTCT): 'Tu thuong' / 'Tu co ket noi'
_g3=wb['3.HTCT_QLVH_G201']; GRP_THUONG,GRP_KN='Tủ thường','Tủ có kết nối'
for _r in range(20,_g3.max_row+1):
    _v=_g3.cell(_r,2).value
    if isinstance(_v,str) and _g3.cell(_r,1).value is None and _v.strip().lower().startswith('tủ'):
        if 'kết nối' in _v.lower(): GRP_KN=_v.strip()
        else: GRP_THUONG=_v.strip()
for _d in qlvh: _d['grp']=GRP_KN if 'kết nối' in (_d['hm'] or '').lower() else GRP_THUONG
qlvh.sort(key=lambda d: 0 if d['grp']==GRP_THUONG else 1)   # nhom Tu thuong truoc (sort on dinh)
hm_order=[]; hm_grp={}; agg={}; byhm=OrderedDict()
for d in qlvh:
    if d['hm'] not in hm_grp: hm_grp[d['hm']]=d['grp']; hm_order.append(d['hm'])
    agg[d['hm']]=agg.get(d['hm'],0)+(d['kl'] or 0)
    byhm.setdefault(d['hm'],OrderedDict()).setdefault(d['phuong'],[]).append(d)
TONG=sum(agg.values())

print('@@Ghi QLVH (2, 3, 8)')
# ================= 3.HTCT_QLVH =================
ws=wb['3.HTCT_QLVH_G201']; ncol=8; tong,sig=markers(ws); first=32
gst=capstyle(ws,32,ncol); dst=capstyle(ws,33,ncol); gh=caph(ws,32); dh=caph(ws,33)
plan=[]; last=None
for h in hm_order:
    if hm_grp[h]!=last: plan.append(('g',h)); last=hm_grp[h]
    plan.append(('d',h))
def w3(ws,fd,pr):
    mg=[]; r=fd; stt=0
    for t,h in plan:
        if t=='g': putstyle(ws,r,gst,ncol); seth(ws,r,gh); ws.cell(r,2,hm_grp[h]); mg.append((r,r,2,3)); r+=1
        else:
            putstyle(ws,r,dst,ncol); seth(ws,r,dh); stt+=1; v=agg[h]
            ws.cell(r,1,stt); ws.cell(r,2,h); ws.cell(r,4,'1 trạm/ngày'); ws.cell(r,5,v); ws.cell(r,6,v); ws.cell(r,7,'=F%d'%r); mg.append((r,r,2,3)); r+=1
    return mg
nt=rewrite(ws,first,tong,ncol,len(plan),w3,33); WROTE['3.HTCT_QLVH_G201']=(first,nt-1)
for cc in 'EFG': ws[cc+str(nt)]='=SUM(%s%d:%s%d)'%(cc,first,cc,nt-1)


# ================= 2.NT_QLVH =================
ws=wb['2.NT_QLVH_G201']; ncol=9; tong,sig=markers(ws); first=44
gst=capstyle(ws,44,ncol); dst=capstyle(ws,45,ncol); gh=caph(ws,44); dh=caph(ws,45)
plan=[]; last=None
for h in hm_order:
    if hm_grp[h]!=last: plan.append(('g',h)); last=hm_grp[h]
    plan.append(('d',h))
def w2(ws,fd,pr):
    mg=[]; r=fd; stt=0; ndata=len(plan)
    for t,h in plan:
        if t=='g': putstyle(ws,r,gst,ncol); seth(ws,r,gh); ws.cell(r,2,hm_grp[h]); mg.append((r,r,2,3)); r+=1
        else:
            putstyle(ws,r,dst,ncol); seth(ws,r,dh); stt+=1; v=agg[h]
            ws.cell(r,1,stt); ws.cell(r,2,h); ws.cell(r,5,v); ws.cell(r,6,'=E%d'%r); ws.cell(r,7,'=E%d'%r); ws.cell(r,8,0); mg.append((r,r,2,3)); r+=1
    ws.cell(fd,4,DBSTR); mg.append((fd,fd+ndata-1,4,4))  # phuong merged whole block
    return mg
nt=rewrite(ws,first,tong,ncol,len(plan),w2,45)
WROTE['2.NT_QLVH_G201']=(first,nt-1)
for cc in 'EFGH': ws[cc+str(nt)]='=SUM(%s%d:%s%d)'%(cc,first,cc,nt-1)


print('@@Ghi bảo dưỡng (6)')
# ================= 6.HTCT_BD =================
ws=wb['6.HTCT_BD_G201']; ncol=8; tong,sig=markers(ws); first=33
gst=capstyle(ws,33,ncol); dst=capstyle(ws,34,ncol); gh=caph(ws,33); dh=caph(ws,34)
def bdg(st): return 'Bảo dưỡng thường xuyên' if st<=95 else 'Bảo dưỡng không thường xuyên'
plan=[]; last=None
for st in sorted(boq):
    gg=bdg(st)
    if gg!=last: plan.append(('g',gg,None)); last=gg
    plan.append(('d',st,boq[st]))
def w6(ws,fd,pr):
    mg=[]; r=fd
    for row in plan:
        if row[0]=='g': putstyle(ws,r,gst,ncol); seth(ws,r,gh); ws.cell(r,2,row[1]); mg.append((r,r,2,3)); r+=1
        else:
            putstyle(ws,r,dst,ncol); seth(ws,r,dh); st=row[1]; it=row[2]
            ws.cell(r,1,st); ws.cell(r,2,it['ten']); ws.cell(r,4,it['dvt']); ws.cell(r,5,it['kh']); ws.cell(r,6,it['th']); ws.cell(r,7,'=F%d'%r); mg.append((r,r,2,3)); r+=1
    return mg
nt=rewrite(ws,first,sig,ncol,len(plan),w6,34); WROTE['6.HTCT_BD_G201']=(first,nt-1)

print('@@Ghi chi tiết Mẫu 4')
# ================= Mau4 (4_G201) QLVH chi tiet =================
ws=wb['4_G201_TTH_PTH_TMY_THU']; ncol=11; tong,sig=markers(ws); first=6
g6=capstyle(ws,6,ncol); h7=capstyle(ws,7,ncol); p8=capstyle(ws,8,ncol); d9=capstyle(ws,9,ncol)
H4=[caph(ws,r) for r in (6,7,8,9)]; MA=ws.cell(9,1).value
plan=[]; last=None
for h in hm_order:
    if hm_grp[h]!=last: plan.append(('grp',hm_grp[h])); last=hm_grp[h]
    plan.append(('hm',h))
    for ph,items in byhm[h].items():
        plan.append(('ph',ph,sum((x['kl'] or 0) for x in items)))
        for x in items: plan.append(('dt',x))
def w4(ws,fd,pr):
    mg=[]; r=fd; hmn=0; hmr=None; phr=None; phs=[]; dstt=0
    def close_ph(end):
        if phr: 
            for L in 'GHIJ': ws[L+str(phr)]='=SUM(%s%d:%s%d)'%(L,phr+1,L,end)
    def close_hm():
        if hmr and phs:
            for L in 'GHIJ': ws[L+str(hmr)]='='+'+'.join(L+str(x) for x in phs)
    for row in plan:
        if row[0] in ('grp','hm'):
            close_ph(r-1); phr=None
        if row[0]=='grp':
            close_hm(); hmr=None; phs=[]
            putstyle(ws,r,g6,ncol); seth(ws,r,H4[0]); ws.cell(r,4,row[1]); mg.append((r,r,4,6)); r+=1
        elif row[0]=='hm':
            close_hm(); phs=[]; dstt=0
            putstyle(ws,r,h7,ncol); seth(ws,r,H4[1]); hmn+=1; h=row[1]; hmr=r
            ws.cell(r,1,MA); ws.cell(r,2,'HM%03d'%hmn); ws.cell(r,3,hmn); ws.cell(r,4,h); mg.append((r,r,4,6)); r+=1
        elif row[0]=='ph':
            close_ph(r-1); phr=r; phs.append(r)
            putstyle(ws,r,p8,ncol); seth(ws,r,H4[2]); ws.cell(r,1,MA); ws.cell(r,2,'HM%03d'%hmn); ws.cell(r,4,'* '+row[1]); mg.append((r,r,4,6)); r+=1
        else:
            x=row[1]; putstyle(ws,r,d9,ncol); seth(ws,r,H4[3]); dstt+=1
            ws.cell(r,1,MA); ws.cell(r,2,'HM%03d'%hmn); ws.cell(r,3,dstt); ws.cell(r,4,x['tu']); ws.cell(r,5,x['duong'] or x['diachi']); ws.cell(r,6,sphuong(x['phuong']))
            ws.cell(r,7,x['kl']); ws.cell(r,8,'=G%d'%r); ws.cell(r,9,'=G%d'%r); ws.cell(r,10,0); r+=1
    close_ph(r-1); close_hm()
    return mg
tail_start = sig
nt=rewrite(ws,first,tail_start,ncol,len(plan),w4,9); WROTE['4_G201_TTH_PTH_TMY_THU']=(first,nt-1)
ws.column_dimensions['A'].hidden=True; ws.column_dimensions['B'].hidden=True

print('@@Ghi chi tiết Mẫu 7')
# ================= Mau7 (7_G201) BD chi tiet (mirror BoQ 499 + breakdown) =================
ws=wb['7_G201_TTH_PTH_TMY_THU']; ncol=13; tong,sig=markers(ws); first=6
g6=capstyle(ws,6,ncol); h7=capstyle(ws,8,ncol); p8=capstyle(ws,9,ncol); d9=capstyle(ws,10,ncol)
H7=[caph(ws,r) for r in (6,8,9,10)]
bqn={norm(v['ten']):k for k,v in boq.items()}
bdhm=OrderedDict()
for d in bd: bdhm.setdefault(norm(d['hm']),OrderedDict()).setdefault(d['phuong'],[]).append(d)
plan=[]; last=None
for st in sorted(boq):
    it=boq[st]; gg=bdg(st)
    if gg!=last: plan.append(('grp',gg)); last=gg
    plan.append(('hm',st,it))
    key=norm(it['ten'])
    if key in bdhm:
        for ph,items in bdhm[key].items():
            plan.append(('ph',ph,sum((x['kl'] or 0) for x in items)))
            for x in items: plan.append(('dt',x))
def w7(ws,fd,pr):
    mg=[]; r=fd; hmr=None; phr=None; phs=[]
    def close_ph(end):
        if phr:
            for L in 'KL': ws[L+str(phr)]='=SUM(%s%d:%s%d)'%(L,phr+1,L,end)
    def close_hm():
        if hmr and phs:
            for L in 'KL': ws[L+str(hmr)]='='+'+'.join(L+str(x) for x in phs)
    for row in plan:
        if row[0] in ('grp','hm'): close_ph(r-1); phr=None; close_hm(); phs=[]; hmr=None
        if row[0]=='grp': putstyle(ws,r,g6,ncol); seth(ws,r,H7[0]); ws.cell(r,4,row[1]); mg.append((r,r,4,9)); r+=1
        elif row[0]=='hm':
            putstyle(ws,r,h7,ncol); seth(ws,r,H7[1]); st=row[1]; it=row[2]; hmr=r
            ws.cell(r,3,st); ws.cell(r,4,it['ten']); ws.cell(r,5,it['dvt']); ws.cell(r,10,it['kh']); ws.cell(r,11,it['th']); ws.cell(r,12,it['nt']); r+=1
        elif row[0]=='ph':
            close_ph(r-1); phr=r; phs.append(r)
            putstyle(ws,r,p8,ncol); seth(ws,r,H7[2]); ws.cell(r,6,'* '+row[1]); mg.append((r,r,6,9)); r+=1
        else:
            x=row[1]; putstyle(ws,r,d9,ncol); seth(ws,r,H7[3])
            ws.cell(r,6,sphuong(x['phuong'])); ws.cell(r,7,x['duong']); ws.cell(r,8,x['tu']); ws.cell(r,9,x['mota']); ws.cell(r,11,x['kl']); ws.cell(r,12,x['kl']); r+=1
    close_ph(r-1); close_hm()
    return mg
nt=rewrite(ws,first,sig,ncol,len(plan),w7,9); WROTE['7_G201_TTH_PTH_TMY_THU']=(first,nt-1)
ws.column_dimensions['A'].hidden=True; ws.column_dimensions['B'].hidden=True


# ================= 8.PYCNT1 (QLVH phieu YCNT) =================
ws=wb['8.PYCNT1_QLVH_G201']; ncol=5; first=33; sig=find_tail(ws,40)
gst=capstyle(ws,33,ncol); dst=capstyle(ws,34,ncol); gh=caph(ws,33); dh=caph(ws,34)
plan=[]; last=None
for h in hm_order:
    if hm_grp[h]!=last: plan.append(('g',h)); last=hm_grp[h]
    plan.append(('d',h))
def w8(ws,fd,pr):
    mg=[]; r=fd; stt=0
    for t,h in plan:
        if t=='g': putstyle(ws,r,gst,ncol); seth(ws,r,gh); ws.cell(r,2,hm_grp[h]); r+=1
        else:
            putstyle(ws,r,dst,ncol); seth(ws,r,dh); stt+=1
            ws.cell(r,1,stt); ws.cell(r,2,h); ws.cell(r,3,'1 trạm/ngày'); ws.cell(r,4,agg[h]); r+=1
    return mg
_nt=rewrite(ws,first,sig,ncol,len(plan),w8,34); WROTE['8.PYCNT1_QLVH_G201']=(first,_nt-1)

# ================= 10.PYCNT2 (BD phieu YCNT) =================
ws=wb['10.PYCNT2_BD_G201']; ncol=5; first=33; sig=find_tail(ws,40)
gst=capstyle(ws,33,ncol); dst=capstyle(ws,34,ncol); gh=caph(ws,33); dh=caph(ws,34)
# BD hang muc co KL, group theo BoQ stt
bqn2={norm(v['ten']):k for k,v in boq.items()}
aggbd={}
for d in bd: aggbd[d['hm']]=aggbd.get(d['hm'],0)+(d['kl'] or 0)
def stof(h):
    k=norm(h); return bqn2.get(k,9999)
plan=[]; last=None
for h in sorted(aggbd,key=stof):
    st=stof(h); gg=bdg(st) if st!=9999 else 'Khác'
    if gg!=last: plan.append(('g',gg,None)); last=gg
    plan.append(('d',h,aggbd[h]))
def w10(ws,fd,pr):
    r=fd; stt=0
    for row in plan:
        if row[0]=='g': putstyle(ws,r,gst,ncol); seth(ws,r,gh); ws.cell(r,2,row[1]); r+=1
        else:
            putstyle(ws,r,dst,ncol); seth(ws,r,dh); stt+=1
            # dvt tu bd
            dvt=next((d['dvt'] for d in bd if d['hm']==row[1]),'')
            ws.cell(r,1,stt); ws.cell(r,2,row[1]); ws.cell(r,3,dvt); ws.cell(r,4,row[2]); r+=1
    return []
_nt=rewrite(ws,first,sig,ncol,len(plan),w10,34); WROTE['10.PYCNT2_BD_G201']=(first,_nt-1)

# ================= 14.VKH_BD + 15 (phat sinh ngoai KH) -> rong =================
# Tim 2 sheet theo TIEU DE (anh co the doi ten 14/15 -> 12/13): khong co thi bo qua
def _hdr(ws,pred):
    for r in range(1,min(ws.max_row,60)+1):
        for c in range(1,14):
            v=ws.cell(r,c).value
            if isinstance(v,str) and pred(v.strip()): return r
for ws in wb.worksheets:
    if ws.sheet_state!='visible': continue
    t=' '.join(str(c.value) for row in ws.iter_rows(max_row=12) for c in row if isinstance(c.value,str)).upper()
    if 'CHI TIẾT' in t and 'NGOÀI KẾ HOẠCH' in t: r0=_hdr(ws,lambda v:v.lower()=='stt'); fd=r0+2 if r0 else None; nc=12
    elif 'PHÁT SINH NGOÀI KẾ HOẠCH' in t: r0=_hdr(ws,lambda v:v=='(1)'); fd=r0+1 if r0 else None; nc=7
    else: continue
    if not fd: print(ws.title,'skip: khong thay dong tieu de bang'); continue
    try:
        rewrite(ws,fd,find_tail(ws,fd),nc,0,lambda a,b,c:[],fd)
    except Exception as e:
        print(ws.title,'skip',e)

def editable(cell):
    return isinstance(cell.value,str) and not isinstance(cell,openpyxl.cell.cell.MergedCell) and cell.value not in PROTECT

print('@@Đổi địa bàn, kỳ, số hiệu')
# ===== DIA BAN: BIA 'Dia ban :' + moi dong 'tren dia ban' (hoa/thuong), tru can cu =====
b=wb['BIA_G201']
for r in range(1,b.max_row+1):
    for c in range(1,5):
        cell=b.cell(r,c)
        if editable(cell) and cell.value.strip().startswith('Địa bàn'): cell.value='Địa bàn : '+DBSTR
enum_pat=re.compile(r'((?:[Pp]hường|[Xx]ã)\s+[^,\n.:]+(?:,\s*(?:[Pp]hường|[Xx]ã)\s+[^,\n.:]+)+)')
for _n in wb.sheetnames:
    ws=wb[_n]
    for row in ws.iter_rows():
        for cell in row:
            if editable(cell) and enum_pat.search(cell.value) and ('trên địa bàn' in cell.value.lower() or cell.value.strip().startswith('Địa bàn')):
                cell.value=enum_pat.sub(DBSTR, cell.value)

# ===== KY: doi thang mau -> thang ky (tru can cu) =====
TM=os.environ.get('THANG_MAU'); TK=os.environ.get('THANG')
if TM and TK and TM!=TK:
    _p=re.compile(r'((?:tháng|THÁNG)\s+)0?%s(?=\s+(?:năm|NĂM)\s+\d{4}|/\d{4})'%TM)
    for _n in wb.sheetnames:
        ws=wb[_n]
        for row in ws.iter_rows():
            for cell in row:
                if editable(cell) and _p.search(cell.value): cell.value=_p.sub(lambda mm: mm.group(1)+TK, cell.value)
    ws=wb['4_G201_TTH_PTH_TMY_THU']
    for r in range(1,ws.max_row+1):
        v=ws.cell(r,1).value
        if isinstance(v,str) and re.fullmatch(r'T\d{2}%02d'%int(TM),v): ws.cell(r,1).value=v[:3]+'%02d'%int(TK)

# ===== DONG PHU THUOC KY trong muc can cu (anh da dong y doi): Bien ban / Nhat ky / Phieu YCNT
#       + SO HIEU van ban (01/YCNT-8-Q7 -> 01/YCNT-7-Q8). Can cu PHAP LY (Nghi dinh, QD, HD, Hop dong, TB, VB) giu nguyen.
QM=os.environ.get('QUAN_MAU','7'); QK=os.environ.get('QUAN','8')
_ky=re.compile(r'^\s*-\s*(Biên bản|Nhật ký|Phiếu yêu cầu)')
_so=re.compile(r'(\d{2}/(?:YCNT|NTCV))-\d{1,2}-Q%s\b'%QM)
_pk=re.compile(r'((?:tháng|THÁNG)\s+)0?%s(?=\s+(?:năm|NĂM)\s+\d{4}|/\d{4})'%(TM or '0'))
for _n in wb.sheetnames:
    ws=wb[_n]
    for row in ws.iter_rows():
        for cell in row:
            v=cell.value
            if not isinstance(v,str) or isinstance(cell,openpyxl.cell.cell.MergedCell): continue
            nv=v
            if v in PROTECT and _ky.match(v):
                if TM and TK: nv=_pk.sub(lambda mm: mm.group(1)+TK, nv)
                if 'trên địa bàn' in nv.lower(): nv=enum_pat.sub(DBSTR, nv)
            if v not in PROTECT or _ky.match(v):
                nv=_so.sub(lambda mm: '%s-%s-Q%s'%(mm.group(1),TK or TM,QK), nv)
            if nv!=v: cell.value=nv

from openpyxl.utils import get_column_letter as _gcl
from openpyxl.cell.cell import MergedCell as _MC
def _cw(ws,c):
    w=ws.column_dimensions[_gcl(c)].width
    return w if w else 8.43
def _span(ws,r,c):
    for m in ws.merged_cells.ranges:
        if m.min_row<=r<=m.max_row and m.min_col<=c<=m.max_col: return m.min_col,m.max_col
    return c,c
def autofit(ws,r1,r2):
    import math
    for r in range(r1,min(r2,ws.max_row)+1):
        need=1
        for c in range(1,ws.max_column+1):
            cell=ws.cell(r,c)
            if isinstance(cell,_MC) or cell.value is None: continue
            t=str(cell.value)
            c1,c2=_span(ws,r,c); w=sum(_cw(ws,cc) for cc in range(c1,c2+1))
            cpl=max(1,int(w*1.05))
            wrap=bool(cell.alignment and cell.alignment.wrap_text)
            parts=t.split(chr(10))
            lines=sum(max(1, math.ceil(len(p)/cpl)) for p in parts) if wrap else len(parts)
            need=max(need,lines)
        cur=ws.row_dimensions[r].height
        if need>1 and cur:
            h=need*15.0
            if cur<h: ws.row_dimensions[r].height=h
for _nm,(_a,_b) in WROTE.items():
    autofit(wb[_nm],_a,_b)
print('@@Lưu file')
wb.calculation.fullCalcOnLoad=True   # Excel tinh lai cong thuc khi mo
wb.save(OUT); print('SAVED',OUT)