# -*- coding: utf-8 -*-
"""Kiem chung tren BO BIEN BAN THAT (13 file, co bien ban 8-11): so hieu = bien ban, KH quy, nhom BD, ngay, vi tri tu.
python3 tests/test_bo_that.py <thu muc PDF/> <mau .xlsm> <bang tra .xlsx> <QUAN> <THANG>"""
import os, sys, re, glob, subprocess, tempfile, json, calendar, warnings
warnings.filterwarnings('ignore')
import openpyxl, pdfplumber
SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'src')
D, MAU, MASTER, QUAN, THANG = sys.argv[1:6]; TMP = tempfile.mkdtemp(prefix='hc_that_')
OK, FAIL = [], []
def check(c, n, d=''): (OK if c else FAIL).append(n + ('' if c else '  <-- ' + str(d)))
out = os.path.join(TMP, 'o.xlsx')
env = dict(os.environ, SRC_DIR=D, MASTER_XLSX=MASTER, SRC_XLSM=MAU, GOI='G201NA', QUAN=QUAN, THANG=THANG, PREFER='', OUT_XLSX=out + '.raw.xlsx')
p = subprocess.run([sys.executable, '-c', "import runpy,report,json;G=runpy.run_path('refill_main.py');R=report.tieout(G);print('CK'+json.dumps(R['checks'],ensure_ascii=False))"], cwd=SRC, env=env, capture_output=True, text=True)
if p.returncode: print(p.stderr[-3000:]); raise SystemExit(1)
ck = json.loads([l for l in p.stdout.splitlines() if l.startswith('CK')][0][2:])
subprocess.run([sys.executable, 'dinhdang_in.py', out + '.raw.xlsx', out], cwd=SRC, check=True, capture_output=True)
for c in ck: check(c['ok'], 'bảng kiểm: ' + c['label'], c['detail'])
wb = openpyxl.load_workbook(out)
T = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
# so hieu that tu bien ban 8-11
that = {}
for f in glob.glob(D + '*.pdf'):
    with pdfplumber.open(f) as pdf: t = ' '.join((pdf.pages[0].extract_text() or '').split())
    for m in re.finditer(r'(\d{2}/\d{2}/(?:YCNT|NTCV)-Q\d+)', t): that[m.group(1)] = 1
mine = set(m.group(1) for s, c, v in T if s[:2] in ('8.', '9.', '10', '11') for m in re.finditer(r'(\d{2}/\d{1,2}/(?:YCNT|NTCV)(?:-Q\d+)?)', v))
check(that and mine == set(that), 'số hiệu sheet 8–11 = biên bản 8–11', (sorted(mine), sorted(that)))
last = calendar.monthrange(2026, int(THANG))[1]
check(not [v for s, c, v in T if re.search(r'ngày\s+(\d+)\s+tháng\s+0?%s\s+năm' % THANG, v, re.I) and int(re.search(r'ngày\s+(\d+)', v, re.I).group(1)) > last], 'không có ngày vượt cuối tháng')
check(any(re.search(r'ngày\s+%d\s+tháng\s+0?%s\s+năm' % (last, THANG), v) for s, c, v in T), 'có ngày %d tháng %s' % (last, THANG))
noft = [w.title for w in wb.worksheets if w.sheet_state == 'visible' and not w.title.startswith('BIA') and 'Trang &P / &N' not in (w.oddFooter.right.text or '')]
check(not noft, 'số trang mọi sheet', noft)
ws6 = wb['6.HTCT_BD_G201']; grp = [ws6.cell(r, 2).value for r in range(33, ws6.max_row) if ws6.cell(r, 1).value is None and ws6.cell(r, 2).value and not str(ws6.cell(r, 2).value).startswith(('Ông', 'Bà'))]
check(len(grp) >= 2, 'sheet 6 có nhóm theo biên bản', grp[:5])
print('\n'.join('ĐẠT  ' + x for x in OK)); print('\n'.join('LỖI  ' + x for x in FAIL)); print('Tổng: %d đạt, %d lỗi' % (len(OK), len(FAIL)))
print('FILE', out)
sys.exit(1 if FAIL else 0)
