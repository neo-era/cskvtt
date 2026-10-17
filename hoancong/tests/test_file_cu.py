# -*- coding: utf-8 -*-
"""Che do 'Bo sung va dinh dang' tren HO SO CU (lap bang ban truoc): mat logo Bia, vung in sheet 4 keo toi cot phu AC,
le cu. Sau dinh dang: le 1,5 cm moi phia, vung in bo cot phu, Bia co lai logo (lay tu mau), PDF co logo.
python3 tests/test_file_cu.py <ho so .xlsx> <mau .xlsm>"""
import os, sys, zipfile, tempfile, subprocess, warnings, re
warnings.filterwarnings('ignore')
H = os.path.dirname(os.path.abspath(__file__)); SRC = os.path.join(H, '..', 'src'); sys.path.insert(0, SRC)
import openpyxl, pdfplumber
HS, MAU = sys.argv[1], sys.argv[2]; T = tempfile.mkdtemp(); OK, FAIL = [], []
def check(c, n, d=''): (OK if c else FAIL).append(n + ('' if c else '  <-- ' + str(d)[:300]))
wb = openpyxl.load_workbook(HS)
bia = [w for w in wb.worksheets if w.title.startswith('BIA')][0]; bia._images = []
s4 = [w for w in wb.worksheets if w.title.startswith('4_')][0]; s4.print_area = 'C1:AC%d' % s4.max_row
for w in wb.worksheets: w.page_margins.top = 0.39; w.page_margins.left = 0.67
cu = os.path.join(T, 'cu.xlsx'); wb.save(cu)
out = os.path.join(T, 'moi.xlsx')
p = subprocess.run([sys.executable, 'dinhdang_in.py', cu, out], cwd=SRC, env=dict(os.environ, SRC_XLSM=MAU), capture_output=True, text=True)
check(p.returncode == 0, 'định dạng chạy được trên hồ sơ cũ', p.stderr[-500:])
w2 = openpyxl.load_workbook(out); cm = 1.5 / 2.54
lech = [(w.title, round(w.page_margins.top * 2.54, 2), round(w.page_margins.bottom * 2.54, 2), round(w.page_margins.left * 2.54, 2), round(w.page_margins.right * 2.54, 2))
        for w in w2.worksheets if w.sheet_state == 'visible' and any(abs(x - cm) > 0.005 for x in (w.page_margins.top, w.page_margins.bottom, w.page_margins.left, w.page_margins.right))]
check(not lech, 'lề trên/dưới/trái/phải 1,5 cm ở mọi sheet', lech[:4])
pa = [w.print_area for w in w2.worksheets if w.title.startswith('4_')][0]
check(re.search(r'\$C\$1:\$K\$', pa or ''), 'sheet 4: bỏ cột phụ AC khỏi vùng in (C:K)', pa)
med = sum(i.file_size for i in zipfile.ZipFile(out).infolist() if '/media/' in i.filename)
check(med > 1000, 'trang Bìa có lại logo (lấy từ file mẫu)', med)
import xuat_pdf as X
R = X.xuat(out, os.path.join(T, 'o.pdf'), os.path.join(H, '..', 'fonts'), log=None)
pg = pdfplumber.open(os.path.join(T, 'o.pdf')).pages
check(len(pg[0].images) >= 1, 'PDF: logo ở trang Bìa')
top = min(c['top'] for c in pg[1].chars); check(abs(top - 1.5 / 2.54 * 72) < 6, 'PDF: chữ đầu trang cách mép trên ~1,5 cm (%.2f cm)' % (top / 72 * 2.54))
print('\n'.join('ĐẠT  ' + x for x in OK)); print('\n'.join('LỖI  ' + x for x in FAIL)); print('Tổng: %d đạt, %d lỗi' % (len(OK), len(FAIL)))
sys.exit(1 if FAIL else 0)
