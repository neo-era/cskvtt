# -*- coding: utf-8 -*-
"""Kiem chung nut "Tai ho so PDF" (src/xuat_pdf.py) - doi chieu voi LibreOffice.
python3 tests/test_pdf.py <ho so .xlsx>
Tieu chi (anh Lam duyet 17/10/2026): xem PDF.md."""
import os, sys, re, time, json, subprocess, tempfile, shutil, collections, warnings
warnings.filterwarnings('ignore')
import openpyxl, pdfplumber
HERE = os.path.dirname(os.path.abspath(__file__)); SRC = os.path.join(HERE, '..', 'src'); sys.path.insert(0, SRC)
FONTS = os.path.join(HERE, '..', 'fonts')
import xuat_pdf as X
XLSX = sys.argv[1]; MAU = sys.argv[2] if len(sys.argv) > 2 else None; TMP = tempfile.mkdtemp(prefix='hc_pdf_')
OK, FAIL = [], []
def check(c, n, d=''): (OK if c else FAIL).append(n + ('' if c else '  <-- ' + str(d)[:400]))

# ---------- don vi: dinh dang so kieu Viet Nam ----------
F = X.dinh_dang
for v, f, k in [(5, '#,##0.000', '5,000'), (1234.5, '#,##0.00', '1.234,50'), (0.09, '0.000', '0,090'), (1234567, '#,##0', '1.234.567'),
                (0.09, 'General', '0,09'), (12, 'General', '12'), (0, '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)', '-'),
                (-5, '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)', '(5)'), (7, '0', '7'), ('abc', '@', 'abc'), (2.5, '#,##0', '3')]:
    check(F(v, f).strip() == k, 'định dạng %r với "%s" → "%s"' % (v, f, k), F(v, f))

# ---------- xuat PDF ----------
out = os.path.join(TMP, 'hs.pdf'); t0 = time.time()
R = X.xuat(XLSX, out, FONTS); dt = time.time() - t0
check(os.path.getsize(out) > 10000, 'tạo được file PDF', R)
check(dt < 30, 'thời gian xuất < 30 giây (CPython: %.1fs)' % dt)
wb = openpyxl.load_workbook(XLSX)
vis = [ws.title for ws in wb.worksheets if ws.sheet_state == 'visible']
check([s for s, n in R['trang']] == vis, 'đúng các sheet đang hiện, đúng thứ tự (bỏ sheet ẩn)', [s for s, n in R['trang']])

# ---------- doi chieu LibreOffice: so trang tung sheet, gia tri cong thuc ----------
def lo_pages(sheet):
    w = openpyxl.load_workbook(XLSX)
    for s in list(w.worksheets):
        if s.title != sheet: w.remove(s)
    w.active = 0
    p = os.path.join(TMP, 'one.xlsx'); w.save(p)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', TMP, p], capture_output=True, timeout=300)
    with pdfplumber.open(os.path.join(TMP, 'one.pdf')) as pdf: return len(pdf.pages)
lo = {s: lo_pages(s) for s in vis}
mine = dict(R['trang'])
lech = {s: (mine.get(s), lo[s]) for s in vis if mine.get(s) != lo[s]}
print('  so trang (PDF, LibreOffice) lech:', lech)
check(all(abs(a - b) <= 2 for a, b in lech.values()), 'số trang từng sheet sát LibreOffice (chênh ≤ 2; LibreOffice co 98,4%% thay vì 100%% như Excel và in tràn ở bảng dài) — %d/%d trang' % (sum(mine.values()), sum(lo.values())), lech)
subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', os.path.join(TMP, 'rc'), XLSX], capture_output=True, timeout=300)
rc = openpyxl.load_workbook(os.path.join(TMP, 'rc', os.path.basename(XLSX)), data_only=True)
ev = X.TinhCongThuc(wb); sai = []; ncf = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                ncf += 1; a = ev.gia_tri(ws.title, c.row, c.column); b = rc[ws.title].cell(c.row, c.column).value
                if not (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 1e-9) and not (a in (None, '', 0) and b in (None, '', 0)):
                    sai.append((ws.title, c.coordinate, a, b))
check(ncf > 0 and not sai, 'giá trị %d ô công thức khớp LibreOffice' % ncf, sai[:5])

# ---------- noi dung: chu cua TUNG O nam dung trong khung o do tren trang PDF ----------
pdf = pdfplumber.open(out); pages = [pg.extract_text() or '' for pg in pdf.pages]
check(sum(n for s, n in R['trang']) == len(pages), 'tổng số trang = cộng các sheet', (sum(n for s, n in R['trang']), len(pages)))
ve = {}                                            # (sheet, r, c) -> [(trang, bbox)]
for gi, (s, vts) in enumerate(R['vi_tri']):
    for (r, c, x, y, w, h) in vts: ve.setdefault((s, r, c), []).append((gi, x, y, w, h))
thieu = []; nkt = 0; chars = {}
for s in vis:
    for (r, c, txt) in X.cac_o_in(wb[s], ev):
        nkt += 1; want = re.sub(r'\s+', '', txt)
        got = ''
        for (gi, x, y, w, h) in ve.get((s, r, c), []):
            if gi not in chars: chars[gi] = pdf.pages[gi].chars
            got += ''.join(ch['text'] for ch in chars[gi] if x - 0.5 <= (ch['x0'] + ch['x1']) / 2 <= x + w + 0.5 and y - 0.5 <= (ch['top'] + ch['bottom']) / 2 <= y + h + 0.5)
        got = re.sub(r'\s+', '', got)
        o = '%s!%s%d' % (s, openpyxl.utils.get_column_letter(c), r)
        if (wb[s].row_dimensions[r].ht or 99) < 10: continue                     # dong dem cua mau (chu an trong mau)
        if want not in got and not any(w.startswith(o + ':') for w in R['canh_bao']): thieu.append((o, want[:30], got[:30]))
check(nkt > 1000 and not thieu, '100%% chữ của %d ô in ra nằm trọn trong khung ô; ô nào bị cắt đều được cảnh báo rõ' % nkt, thieu[:6])
print('  cảnh báo cắt chữ:', [w for w in R['canh_bao'] if 'bị cắt' in w])

# ---------- chan trang, kho giay, tieng Viet ----------
start = 0; bad = []
for s, n in R['trang']:
    for i in range(n):
        t = pages[start + i]; pg = pdf.pages[start + i]
        if not s.startswith('BIA') and 'Trang %d / %d' % (i + 1, n) not in t: bad.append((s, i + 1, t[-40:]))
        ori = wb[s].page_setup.orientation or 'portrait'
        if (pg.width > pg.height) != (ori == 'landscape') or abs(max(pg.width, pg.height) - 841.89) > 1: bad.append((s, 'khổ', pg.width, pg.height))
    start += n
check(not bad, 'mỗi trang có "Trang x / y" theo từng biên bản + đúng khổ A4 ngang/dọc', bad[:4])
al = '\n'.join(pages)
check('Độc lập - Tự do - Hạnh phúc' in al and re.search(r'ngày 30 tháng 9 năm 2026', al), 'tiếng Việt đủ dấu, chọn/tìm được chữ')
check(any('5,000' in p or ',000' in p for p in pages), 'số hiện kiểu Việt Nam trong PDF')

# ---------- khong con o bi cat chu (chieu cao dong du cho so dong Excel ngat) ----------
wbr = openpyxl.load_workbook(XLSX, rich_text=True); evr = X.TinhCongThuc(wbr); D = X.PDF(FONTS); cat = []
for ws in wbr.worksheets:
    if ws.sheet_state != 'visible': continue
    S = X.Sheet(ws, evr, D, X._theme(wbr), X._mdw(wbr, D))
    for r in S.rows:
        ht = ws.row_dimensions[r].ht
        if ht is None or ht < 10: continue
        for c in S.cols:
            a = S.mgcell.get((r, c))
            if a and (a != (r, c) or S.mg[a][0] > r): continue
            x = ws.cell(r, c)
            if x.value is None or not x.alignment.wrap_text: continue
            runs, so = S.chu(x)
            if so or not runs or not ''.join(t[0] for t in runs).strip(): continue
            need = sum(l[1] for l in S.ngat(runs, S.rong_o(r, c) - 2 * X.PAD - S._lui(x)))
            if need > ht + 0.5: cat.append((ws.title, x.coordinate, ht, round(need, 1)))
check(not cat, 'không còn ô bị cắt chữ (chiều cao dòng đủ cho số dòng ngắt)', cat[:5])

# ---------- doi chieu cach ngat dong voi chinh Excel: cac dong Excel tu canh chieu cao trong file mau ----------
if MAU:
    import zipfile
    z = zipfile.ZipFile(MAU); wm = openpyxl.load_workbook(MAU, rich_text=True); evm = X.TinhCongThuc(wm); Dm = X.PDF(FONTS)
    rel = z.read('xl/_rels/workbook.xml.rels').decode('utf8'); wx = z.read('xl/workbook.xml').decode('utf8')
    tg = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship [^>]*?Id="(rId\d+)"[^>]*?Target="(?:/xl/)?worksheets/(sheet\d+\.xml)"', rel)}
    tg.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship [^>]*?Target="(?:/xl/)?worksheets/(sheet\d+\.xml)"[^>]*?Id="(rId\d+)"', rel)})
    ten = {m.group(2): m.group(1) for m in re.finditer(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="(rId\d+)"', wx)}
    khop = tong = 0
    for rid, f in tg.items():
        if rid not in ten: continue
        ws = wm[ten[rid].replace('&amp;', '&')]; S = X.Sheet(ws, evm, Dm, X._theme(wm), X._mdw(wm, Dm))
        for rr, at in re.findall(r'<row r="(\d+)"([^>]*)>', z.read('xl/worksheets/' + f).decode('utf8')):
            if 'ht=' not in at or 'customHeight' in at: continue
            r = int(rr); ht = float(re.search(r'ht="([\d.]+)"', at).group(1)); best = 0; co = False
            for c in range(1, ws.max_column + 1):
                a = S.mgcell.get((r, c))
                if a and (a != (r, c) or S.mg[a][0] > r): continue
                x = ws.cell(r, c)
                if x.value is None or not x.alignment.wrap_text: continue
                runs, so = S.chu(x)
                if so or not runs: continue
                co = True; best = max(best, sum(l[1] for l in S.ngat(runs, S.rong_o(r, c) - 2 * X.PAD - S._lui(x))))
            if co: tong += 1; khop += abs(best - ht) < 0.6
    check(tong > 500 and khop / tong >= 0.99, 'ngắt dòng khớp Excel: %d/%d dòng Excel tự canh chiều cao (%.1f%%)' % (khop, tong, 100.0 * khop / max(1, tong)))

# ---------- ca xau ----------
w2 = openpyxl.load_workbook(XLSX); s10 = [s for s in w2.sheetnames if s.startswith('10')][0]; ws = w2[s10]
ws['B34'] = '=VLOOKUP(1,A1:B2,2,0)'
ws['B35'] = ('Ô rất dài ' * 600).strip()
ws.row_dimensions[36].hidden = True; an = ws['B36'].value
ws.print_area = None
p2 = os.path.join(TMP, 'xau.xlsx'); w2.save(p2); o2 = os.path.join(TMP, 'xau.pdf')
try:
    R2 = X.xuat(p2, o2, FONTS); ok = True
except Exception as e:
    ok = False; R2 = {'canh_bao': [str(e)]}
check(ok, 'ca xấu: ô quá dài + công thức lạ + không có vùng in vẫn xuất được', R2.get('canh_bao'))
check(ok and any(s10 in w and 'B34' in w for w in R2['canh_bao']), 'ca xấu: báo rõ ô có công thức không tính được (%s!B34)' % s10, R2.get('canh_bao'))
if ok:
    t2 = ' '.join((pg.extract_text() or '') for pg in pdfplumber.open(o2).pages)
    st = 0; tx = ''
    for s, n in R2['trang']:
        if s == s10: tx = ' '.join((pg.extract_text() or '') for pg in pdfplumber.open(o2).pages[st:st + n])
        st += n
    check(isinstance(an, str) and an not in tx, 'ca xấu: dòng ẩn không in ra', an)

print('\n'.join('ĐẠT  ' + x for x in OK)); print('\n'.join('LỖI  ' + x for x in FAIL))
print('Tổng: %d đạt, %d lỗi  | %s  %.1fs' % (len(OK), len(FAIL), out, dt))
sys.exit(1 if FAIL else 0)
