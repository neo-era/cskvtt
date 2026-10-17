# -*- coding: utf-8 -*-
"""Kiem chung tieu chi dot sua 10/2026 (xem DESIGN.md).
python3 tests/test_v2.py <thu muc PDF (co / cuoi)> <mau .xlsm> <bang tra .xlsx>
Can: bo bien ban T08 Q8 (co bien ban so 5 QLVH tong hop). Chay ky 9 va ky 2 tren cung bo PDF."""
import os, sys, re, subprocess, tempfile, shutil, glob, calendar, warnings
warnings.filterwarnings('ignore')
import openpyxl, pdfplumber
HERE = os.path.dirname(os.path.abspath(__file__)); SRC = os.path.join(HERE, '..', 'src')
PDF_DIR, MAU, MASTER = sys.argv[1], sys.argv[2], sys.argv[3]
TMP = tempfile.mkdtemp(prefix='hc_test_')
OK = []; FAIL = []
def check(cond, name, detail=''):
    (OK if cond else FAIL).append(name + ('' if cond else '  <-- ' + str(detail)))

def run(thang, mau, out, pdf_dir=PDF_DIR):
    env = dict(os.environ, SRC_DIR=pdf_dir, MASTER_XLSX=MASTER, SRC_XLSM=mau, GOI='G201NA', QUAN='8', THANG=str(thang),
               PREFER='Phường Chánh Hưng;Phường Phú Định;Phường Bình Đông', OUT_XLSX=out + '.raw.xlsx')
    for k in ('THANG_MAU', 'QUAN_MAU'): env.pop(k, None)
    code = ("import runpy,report,json;G=runpy.run_path('refill_main.py');R=report.tieout(G);"
            "print('CHECKS',json.dumps([[c['ok'],c['label']] for c in R['checks']],ensure_ascii=False))")
    p = subprocess.run([sys.executable, '-c', code], cwd=SRC, env=env, capture_output=True, text=True)
    if p.returncode: print(p.stdout[-2000:], p.stderr[-3000:]); raise SystemExit('pipeline loi')
    checks = [l for l in p.stdout.splitlines() if l.startswith('CHECKS')][0]
    subprocess.run([sys.executable, 'dinhdang_in.py', out + '.raw.xlsx', out], cwd=SRC, env=env, check=True, capture_output=True)
    return checks

def texts(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str): yield ws.title, c.coordinate, c.value

def find_pdf(pred):
    for p in sorted(glob.glob(PDF_DIR + '*.pdf')):
        with pdfplumber.open(p) as pdf:
            tb = [r for pg in pdf.pages[:3] for t in pg.extract_tables() for r in t]
            if pred(pdf, tb): return p

# ---------- ky 9 ----------
out9 = os.path.join(TMP, 't9.xlsx'); ck9 = run(9, MAU, out9)
wb = openpyxl.load_workbook(out9); mau = openpyxl.load_workbook(MAU)
allt = list(texts(wb))
# (1) ngay cuoi thang
bad = [(s, c, v[:60]) for s, c, v in allt if re.search(r'ngày\s+31\s+tháng\s+0?9\s+năm', v)]
check(not bad, '(1) không còn "ngày 31 tháng 9"', bad[:3])
good = [v for s, c, v in allt if re.search(r'ngày\s+30\s+tháng\s+0?9\s+năm\s+2026', v)]
check(len(good) >= 10, '(1) có "ngày 30 tháng 9 năm 2026" ở các sheet', len(good))
check(any(re.search(r'ngày\s+01\s+tháng\s+0?9', v) for s, c, v in allt), '(1) ngày 01 giữ nguyên')
# (2) KH quy sheet 3 = bien ban 5
p5 = find_pdf(lambda pdf, tb: any('kế hoạch quý' in ' '.join(' '.join((x or '') for x in r).split()).lower() for r in tb) and any('Duy trì trạm' in (r[1] or '') for r in tb if len(r) > 1))
check(p5 is not None, '(2) có biên bản số 5 trong bộ test')
if p5 is None: raise SystemExit('\n'.join(FAIL))
kh = []
with pdfplumber.open(p5) as pdf:
    for pg in pdf.pages:
        for t in pg.extract_tables():
            for r in t:
                if len(r) >= 6 and re.fullmatch(r'\d+', (r[0] or '').strip()) and 'Duy trì' in (r[1] or ''):
                    kh.append(float((r[3] or '0').replace('.', '').replace(',', '.')))
ws3 = wb['3.HTCT_QLVH_G201']
e3 = [ws3.cell(r, 5).value for r in range(32, ws3.max_row) if isinstance(ws3.cell(r, 1).value, int)]
check(sorted(e3) == sorted(kh) and len(kh) > 0, '(2) sheet 3 cột E = KH quý biên bản 5', (sorted(e3)[:5], sorted(kh)[:5]))
f3 = [ws3.cell(r, 6).value for r in range(32, ws3.max_row) if isinstance(ws3.cell(r, 1).value, int)]
check(e3 != f3, '(2) cột E khác cột F (không phải số tháng)')
# (3) sheet 4 vi tri dat tu = dia chi PDF4
p4 = find_pdf(lambda pdf, tb: sum(1 for r in tb if len(r) == 9) >= 5)
addr = {}
with pdfplumber.open(p4) as pdf:
    for pg in pdf.pages[:3]:
        for t in pg.extract_tables():
            for r in t:
                r = [(c or '').replace('\n', ' ').strip() for c in r]
                if len(r) >= 9 and r[3] == '8' and r[2]: addr[r[2]] = 1
ws4 = wb['4_G201_TTH_PTH_TMY_THU']
e4 = set(ws4.cell(r, 5).value for r in range(6, ws4.max_row) if isinstance(ws4.cell(r, 3).value, int) and ws4.cell(r, 7).value not in (None,) and not str(ws4.cell(r, 7).value).startswith('='))
hit = [a for a in addr if a in e4]
check(len(addr) > 5 and len(hit) == len(addr), '(3) sheet 4 cột E = địa chỉ đặt tủ trong biên bản', (len(hit), len(addr), list(addr)[:3]))
# (5) so trang
noft = [s.title for s in wb.worksheets if s.sheet_state == 'visible' and not s.title.startswith('BIA') and 'Trang &P / &N' not in (s.oddFooter.right.text or '')]
check(not noft, '(5) chân trang "Trang &P / &N" mọi sheet (trừ Bìa)', noft)
# chung: can cu giu nguyen
legal = re.compile(r'^\s*-?\s*(Nghị định|Quyết định|Hướng dẫn|Căn cứ|Thông báo|Văn bản số)')
mt = set(v for s, c, v in texts(mau) if legal.match(v)); ot = set(v for s, c, v in allt)
miss = [v[:70] for v in mt if v not in ot]
check(not miss, 'chung: căn cứ pháp lý giữ nguyên', miss[:2])
check(all(x[0] for x in __import__('json').loads(ck9[7:]) if 'Không còn' not in x[1]), 'chung: bảng kiểm report đạt', ck9)

# ---------- (4) so hieu: mau cu (-M-Qn) va mau moi (STT/Thang/YCNT) ----------
def sohieu(wb):
    return [(s, c, v) for s, c, v in texts(wb) if re.search(r'(YCNT|NTCV)', v) and s[:2] in ('8.', '9.', '10', '11')]
sh = sohieu(wb)
check(sh and all(re.search(r'\d{2}/0?9/(YCNT|NTCV)', v) for s, c, v in sh) and not any(re.search(r'-Q\d', v) for s, c, v in sh),
      '(4) mẫu cũ → STT/Tháng/YCNT tháng 9, bỏ quận', [v[-40:] for s, c, v in sh][:4])
# mau moi: 01/08/YCNT (2 chu so) -> 01/09/YCNT
m2 = openpyxl.load_workbook(MAU, keep_vba=True)
for ws in m2.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and re.search(r'\d{2}/(YCNT|NTCV)-\d{1,2}-Q\d+', c.value):
                c.value = re.sub(r'(\d{2})/(YCNT|NTCV)-(\d{1,2})-Q\d+', lambda m: '%s/%02d/%s' % (m.group(1), int(m.group(3)), m.group(2)), c.value)
mau2 = os.path.join(TMP, 'mau_moi.xlsm'); m2.save(mau2)
out9b = os.path.join(TMP, 't9b.xlsx'); run(9, mau2, out9b)
sh = sohieu(openpyxl.load_workbook(out9b))
check(sh and all(re.search(r'\d{2}/09/(YCNT|NTCV)', v) for s, c, v in sh), '(4) mẫu mới 01/08/YCNT → 01/09/YCNT (giữ 2 chữ số)', [v[-40:] for s, c, v in sh][:4])

# ---------- ky 2 (28 ngay) ----------
out2 = os.path.join(TMP, 't2.xlsx'); run(2, MAU, out2)
t2 = [v for s, c, v in texts(openpyxl.load_workbook(out2))]
check(any(re.search(r'ngày\s+28\s+tháng\s+0?2\s+năm', v) for v in t2) and not any(re.search(r'ngày\s+3[01]\s+tháng\s+0?2\s', v) for v in t2), '(1) kỳ tháng 2 → ngày 28')

# ---------- thieu bien ban 5 ----------
d5 = os.path.join(TMP, 'pdf_no5') + '/'; os.makedirs(d5)
for p in glob.glob(PDF_DIR + '*.pdf'):
    if p != p5: shutil.copy(p, d5)
out_n5 = os.path.join(TMP, 'n5.xlsx'); ckn5 = run(9, MAU, out_n5, d5)
ws3 = openpyxl.load_workbook(out_n5)['3.HTCT_QLVH_G201']
e3 = [ws3.cell(r, 5).value for r in range(32, ws3.max_row) if isinstance(ws3.cell(r, 1).value, int)]
check(all(v is None for v in e3) and 'false, "Kế hoạch quý' in ckn5.replace("[false", "false"), 'ca xấu: thiếu biên bản 5 → cột E trống + mục kiểm Xem lại', (e3[:3], ckn5[:300]))

# ---------- 0 loi cong thuc (LibreOffice) ----------
rc = os.path.join(TMP, 'rc'); subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', rc, out9], capture_output=True, timeout=600)
errs = [(s, c) for s, c, v in texts(openpyxl.load_workbook(os.path.join(rc, 't9.xlsx'), data_only=True)) if v.strip().upper() in ('#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A', '#NUM!', '#NULL!')]
check(not errs, 'chung: 0 lỗi công thức', errs[:3])

print('\n'.join('ĐẠT  ' + x for x in OK)); print('\n'.join('LỖI  ' + x for x in FAIL))
print('Tổng: %d đạt, %d lỗi' % (len(OK), len(FAIL)))
sys.exit(1 if FAIL else 0)
